"""
Script to generate the Authentic Work Experience Presentation
Candidate: Maheshchandra Hegde
Profile: Product Design Expert | UX Architect | Creative Technologist (16+ Years Experience)
Position applied: Senior Manager - IT
Format: 16:9 Widescreen PowerPoint Presentation (.pptx)
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_pptx_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Elegant, Authentic Executive Palette
    COLOR_PRIMARY_DARK = RGBColor(15, 23, 42)     # Slate 900 #0F172A
    COLOR_PRIMARY_BLUE = RGBColor(2, 132, 199)    # Sky / Tech Blue #0284C7
    COLOR_BRAND_DEEP   = RGBColor(14, 116, 144)   # Deep Ocean Teal #0E7490
    COLOR_ACCENT_AMBER = RGBColor(217, 119, 6)    # Warm Amber #D97706
    COLOR_BG_LIGHT     = RGBColor(248, 250, 252)  # Canvas Light #F8FAFC
    COLOR_CARD_BG      = RGBColor(255, 255, 255)  # Pure White #FFFFFF
    COLOR_CARD_BORDER  = RGBColor(226, 232, 240)  # Border Subtle #E2E8F0
    COLOR_TEXT_MAIN    = RGBColor(30, 41, 59)     # Slate 800 #1E293B
    COLOR_TEXT_MUTED   = RGBColor(100, 116, 139)  # Slate 500 #64748B
    COLOR_SUCCESS      = RGBColor(16, 149, 99)    # Emerald Green #109563
    COLOR_CYAN_ACCENT  = RGBColor(56, 189, 248)   # Cyan #38BDF8

    TOTAL_SLIDES = 11

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, category, title, dark=False):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_CYAN_ACCENT if dark else COLOR_BRAND_DEEP

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.55))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255) if dark else COLOR_PRIMARY_DARK

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.025))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(51, 65, 85) if dark else COLOR_CARD_BORDER
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    def add_footer(slide, current_page, total_pages=TOTAL_SLIDES, dark=False):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
        tf = footer_box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Maheshchandra Hegde · Technical Work Experience & Capability Presentation | Slide {current_page} of {total_pages}"
        p.font.size = Pt(9)
        p.font.color.rgb = RGBColor(148, 163, 184) if dark else COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 1: Title Slide (Executive Dark)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, COLOR_PRIMARY_DARK)

    # Accent decorative bar
    strip = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(0.25), Inches(4.8))
    strip.fill.solid()
    strip.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
    strip.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.35), Inches(1.3), Inches(11.2), Inches(3.4))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "TECHNICAL LEADERSHIP & WORK EXPERIENCE PRESENTATION"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_CYAN_ACCENT
    p0.space_after = Pt(10)

    p1 = tf1.add_paragraph()
    p1.text = "Maheshchandra Hegde"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_after = Pt(4)

    p2 = tf1.add_paragraph()
    p2.text = "Founder & CTO | Building AI‑Driven 3D Asset Platforms & Spatial Computing Solutions @srushtilabs.com"
    p2.font.size = Pt(13.5)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_CYAN_ACCENT
    p2.space_after = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "16+ years delivering scalable enterprise frontend systems, high-reliability clinical platforms, offline-first architectures, and generative 3D spatial tools.\nBangalore, India | hid.mahesh@gmail.com | +91 9535253329 / 7022407280 | hegdemahesh.in | linkedin.com/in/maheshchandrahegde"
    p3.font.size = Pt(11)
    p3.font.color.rgb = RGBColor(148, 163, 184)

    # 4 Authentic Pillar Cards
    pillars = [
        ("Scalable Enterprise Systems", "Philips Healthcare (Cyient), Cisco Systems, enterprise web apps, micro-frontends"),
        ("Startup Founder & Tech Director", "Technoyana (technoyana.in), Twitan.com (high-reliability sports OS), InnoBrik"),
        ("3D Spatial & Generative AI", "Voxelforge AI (srushtilabs.com/voxelforge), ayam3d, CAE flight simulation"),
        ("Architectural Heritage", "Rooted in CAD & 3D visualization assisting architect father; 50+ 3D projects delivered")
    ]
    card_w = Inches(2.75)
    card_h = Inches(1.35)
    start_x = Inches(1.35)
    start_y = Inches(5.0)
    gap = Inches(0.24)

    for i, (title, desc) in enumerate(pillars):
        cx = start_x + i * (card_w + gap)
        add_card(s1, cx, start_y, card_w, card_h, bg_color=RGBColor(24, 34, 53), border_color=RGBColor(51, 65, 85))
        tb = s1.shapes.add_textbox(cx + Inches(0.15), start_y + Inches(0.12), card_w - Inches(0.3), card_h - Inches(0.24))
        ctf = tb.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(11.5)
        cp1.font.bold = True
        cp1.font.color.rgb = RGBColor(255, 255, 255)
        cp1.space_after = Pt(4)
        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(9)
        cp2.font.color.rgb = RGBColor(148, 163, 184)

    add_footer(s1, 1, TOTAL_SLIDES, dark=True)

    # ==========================================
    # SLIDE 2: Professional Profile & Architectural Roots
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, COLOR_BG_LIGHT)
    add_header(s2, "Executive Profile", "Professional Background, Architectural Heritage & Technical Philosophy")

    # Left: Heritage & Story Card (Width: 5.6)
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    hb = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.81), Inches(1.61), Inches(5.58), Inches(0.42))
    hb.fill.solid()
    hb.fill.fore_color.rgb = COLOR_BRAND_DEEP
    hb.line.fill.background()

    ht = s2.shapes.add_textbox(Inches(0.95), Inches(1.68), Inches(5.3), Inches(0.35))
    ht.text_frame.paragraphs[0].text = "ARCHITECTURAL LINEAGE & PASSION FOR 3D"
    ht.text_frame.paragraphs[0].font.size = Pt(11)
    ht.text_frame.paragraphs[0].font.bold = True
    ht.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    s_tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(5.2), Inches(4.5))
    stf = s_tb.text_frame
    stf.word_wrap = True

    p = stf.paragraphs[0]
    p.text = "Natural Convergence of Architecture and Computing"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK
    p.space_after = Pt(6)

    story_text = (
        "• Early Foundation with Architect Father: Growing up assisting my architect father in his studio provided my earliest foundation in drafting, spatial coordination, blueprints, CAD, and architectural drawings. This practical exposure sparked an enduring passion for 3D modeling, spatial thinking, and visualization.\n\n"
        "• Blending Engineering with Human Interface: Pursued formal education with a Bachelor of Engineering in Electronics & Communication and a Master of Science in Computing (Robert Gordon University, UK), specializing in Human Interface Design & Development.\n\n"
        "• 16+ Years Hands-on Leadership: Combined rigorous software architecture with creative spatial technologies—spanning flight simulation visual databases (CAE), enterprise dashboards (Cisco), healthcare platforms (Philips), startup products (Technoyana, Twitan), and generative 3D AI (Voxelforge AI, ayam3d)."
    )
    p2 = stf.add_paragraph()
    p2.text = story_text
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_TEXT_MAIN

    # Right: Career Snapshot & Core Competencies (Width: 5.85)
    rx = Inches(6.68)
    rw = Inches(5.85)

    add_card(s2, rx, Inches(1.6), rw, Inches(5.2))
    rb = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, rx + Inches(0.01), Inches(1.61), rw - Inches(0.02), Inches(0.42))
    rb.fill.solid()
    rb.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
    rb.line.fill.background()

    rt = s2.shapes.add_textbox(rx + Inches(0.15), Inches(1.68), rw - Inches(0.3), Inches(0.35))
    rt.text_frame.paragraphs[0].text = "CAREER SNAPSHOT & CORE COMPETENCIES"
    rt.text_frame.paragraphs[0].font.size = Pt(11)
    rt.text_frame.paragraphs[0].font.bold = True
    rt.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    rtb = s2.shapes.add_textbox(rx + Inches(0.2), Inches(2.15), rw - Inches(0.4), Inches(4.5))
    rtf = rtb.text_frame
    rtf.word_wrap = True

    c_items = [
        ("Frontend & Scalable Architecture", "React, Angular, TypeScript, Node.js, Web Components, Micro-frontends, LitElement, Design Systems."),
        ("Cloud, Edge & Reliability", "GCP, Firebase, AWS, offline-first architectures, IndexedDB/Service Workers, WebSockets, real-time data sync."),
        ("Spatial 3D & Generative AI", "AI-assisted 3D generation (Voxelforge AI), automated mesh retopology (ayam3d), WebGL, Three.js, CAD drafting, 3D walkthroughs."),
        ("Product & Startup Leadership", "Founder/Co-Founder across Technoyana, Twitan.com, SrushtiLabs, InnoBrik; product strategy, sprint governance, mentoring."),
        ("Key Industry Tenures", "Cyient (Philips Healthcare), UST Global (Cisco Systems), Ness Technologies, Moonraft Innovation Labs, ThoughtFocus, CAE Simulation.")
    ]
    for ci_i, (clbl, cval) in enumerate(c_items):
        cp1 = rtf.paragraphs[0] if ci_i == 0 else rtf.add_paragraph()
        cp1.text = f"• {clbl}: "
        cp1.font.bold = True
        cp1.font.size = Pt(10)
        cp1.font.color.rgb = COLOR_PRIMARY_DARK
        cp2 = rtf.add_paragraph()
        cp2.text = f"  {cval}"
        cp2.font.size = Pt(9)
        cp2.font.color.rgb = COLOR_TEXT_MUTED
        cp2.space_after = Pt(4)

    add_footer(s2, 2, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 3: Project 1 — Philips ICCA (Cyient)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, COLOR_BG_LIGHT)
    add_header(s3, "Enterprise Project Experience 1", "Healthcare Mission-Critical Systems: Philips ICCA (Cyient Limited)")

    # Left: Role & Context Card
    add_card(s3, Inches(0.8), Inches(1.6), Inches(4.3), Inches(5.2))
    p1_tb = s3.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(3.8), Inches(4.8))
    p1_tf = p1_tb.text_frame
    p1_tf.word_wrap = True

    p = p1_tf.paragraphs[0]
    p.text = "Philips ICCA"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK
    p.space_after = Pt(4)

    p = p1_tf.add_paragraph()
    p.text = "Intellispace Critical Care & Anesthesia"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_BRAND_DEEP
    p.space_after = Pt(10)

    p1_meta = [
        ("Role", "Senior Technical Lead (React / NodeJS)"),
        ("Tenure", "Nov 2023 – Mar 2025"),
        ("Company", "Cyient Limited for Philips Healthcare, Bangalore"),
        ("Operational Reality", "24/7 continuous ICU hospital environments; zero tolerance for UI freeze or missing patient vitals"),
        ("Key Focus", "Scalable frontend architecture, high performance, strict healthcare standards compliance")
    ]
    for lbl, val in p1_meta:
        p = p1_tf.add_paragraph()
        p.text = f"{lbl}: "
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_PRIMARY_DARK
        p2 = p1_tf.add_paragraph()
        p2.text = f"{val}"
        p2.font.size = Pt(9)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(5)

    # Right: Objectives & Technical Delivery
    r3_x = Inches(5.35)
    r3_w = Inches(7.18)

    p1_cards = [
        ("1. High-Performance Clinical Telemetry & Zero Latency",
         "Architected React-based frontend modules responsible for rendering complex clinical time-series charts, vitals, and medication administration workflows. Enforced strict memory management and component memoization to prevent garbage collection pauses during long-running bedside sessions."),
        ("2. Decoupled Modular Architecture & Maintainability",
         "Deconstructed complex clinical workflows into reusable, modular component libraries. Created robust contract boundaries with backend REST and streaming APIs, accelerating multi-developer collaboration while eliminating regression risks across unrelated clinical modules."),
        ("3. Healthcare Regulatory & Security Standards Compliance",
         "Ensured all developed UI modules strictly complied with medical software safety benchmarks, HIPAA data protection guidelines, role-based clinician access controls, and robust audit trails for all patient data modifications.")
    ]

    for i, (title, desc) in enumerate(p1_cards):
        cy = Inches(1.6) + i * Inches(1.76)
        add_card(s3, r3_x, cy, r3_w, Inches(1.6))
        otb = s3.shapes.add_textbox(r3_x + Inches(0.2), cy + Inches(0.15), r3_w - Inches(0.4), Inches(1.3))
        otf = otb.text_frame
        otf.word_wrap = True
        op1 = otf.paragraphs[0]
        op1.text = title
        op1.font.size = Pt(12)
        op1.font.bold = True
        op1.font.color.rgb = COLOR_PRIMARY_BLUE
        op1.space_after = Pt(5)
        op2 = otf.add_paragraph()
        op2.text = desc
        op2.font.size = Pt(9.5)
        op2.font.color.rgb = COLOR_TEXT_MAIN

    add_footer(s3, 3, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 4: Project 2 — Startup Ventures: Technoyana & Twitan.com
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, COLOR_BG_LIGHT)
    add_header(s4, "Startup Ventures & Cloud Systems", "Product Engineering & Edge Resilience: Technoyana & Twitan.com")

    # Left: Startup Snapshot Card
    add_card(s4, Inches(0.8), Inches(1.6), Inches(4.3), Inches(5.2))
    p2_tb = s4.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(3.8), Inches(4.8))
    p2_tf = p2_tb.text_frame
    p2_tf.word_wrap = True

    p = p2_tf.paragraphs[0]
    p.text = "Technoyana & Twitan"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK
    p.space_after = Pt(4)

    p = p2_tf.add_paragraph()
    p.text = "technoyana.in | twitan.com"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_BLUE
    p.space_after = Pt(10)

    p2_meta = [
        ("Technoyana (2021 – 2023)", "Co-Founder & Director; delivered fintech apps, product discovery engines, web & cross-platform mobile apps (React Native / Ionic)."),
        ("Twitan.com (Apr 2025 – Present)", "Founder & Product Design Lead; AI-driven sports management, bracket engines & live scoring (Shutlify badminton OS)."),
        ("Operational Reality", "Tournament venues with dead-zones, unstable cellular Wi-Fi, and real-time live match arbitration."),
        ("Tech Stack", "React, Angular, Node.js, Firebase/GCP, PWA, IndexedDB, WebSockets.")
    ]
    for lbl, val in p2_meta:
        p = p2_tf.add_paragraph()
        p.text = f"{lbl}: "
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_PRIMARY_DARK
        p2 = p2_tf.add_paragraph()
        p2.text = f"{val}"
        p2.font.size = Pt(9)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(5)

    # Right: 3 Core Architectural Achievements
    r4_x = Inches(5.35)
    r4_w = Inches(7.18)

    p2_cards = [
        ("1. Offline-First Resilience for Unstable Physical Venues",
         "Architected an offline-first state machine using IndexedDB local storage and Service Workers for courtside officials. When Wi-Fi or cellular connections fluctuate, the app operates without interruption; all scores and events queue locally and sync with conflict-resolution algorithms upon reconnection."),
        ("2. Cloud Architecture, Firebase & Rapid Delivery Workflows",
         "Set up scalable backend cloud services using GCP, Firebase, and Node.js microservices. Evaluated and optimized tech stacks for clients, drastically reducing delivery timelines for complex fintech and consumer applications while keeping cloud operating expenses lean."),
        ("3. Modular Athlete & Tournament Dashboards",
         "Designed modular dashboards integrating live statistics, tournament bracket generation (knockout, round-robin), court scheduling engines, and analytical insights, recognized for exceptional user experience and responsiveness.")
    ]

    for i, (title, desc) in enumerate(p2_cards):
        cy = Inches(1.6) + i * Inches(1.76)
        add_card(s4, r4_x, cy, r4_w, Inches(1.6))
        otb = s4.shapes.add_textbox(r4_x + Inches(0.2), cy + Inches(0.15), r4_w - Inches(0.4), Inches(1.3))
        otf = otb.text_frame
        otf.word_wrap = True
        op1 = otf.paragraphs[0]
        op1.text = title
        op1.font.size = Pt(12)
        op1.font.bold = True
        op1.font.color.rgb = COLOR_PRIMARY_BLUE
        op1.space_after = Pt(5)
        op2 = otf.add_paragraph()
        op2.text = desc
        op2.font.size = Pt(9.5)
        op2.font.color.rgb = COLOR_TEXT_MAIN

    add_footer(s4, 4, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 5: Project 3 — Spatial AI & Generative 3D (Voxelforge AI & ayam3d)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, COLOR_BG_LIGHT)
    add_header(s5, "Generative AI & Spatial Tech", "Spatial AI & Generative 3D Pipelines: SrushtiLabs (Voxelforge AI) & ayam3d")

    # Left: SrushtiLabs & Voxelforge AI Card
    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    hb5 = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.81), Inches(1.61), Inches(5.58), Inches(0.42))
    hb5.fill.solid()
    hb5.fill.fore_color.rgb = COLOR_BRAND_DEEP
    hb5.line.fill.background()

    ht5 = s5.shapes.add_textbox(Inches(0.95), Inches(1.68), Inches(5.3), Inches(0.35))
    ht5.text_frame.paragraphs[0].text = "SRUSHTILABS · VOXELFORGE AI (srushtilabs.com/voxelforge/)"
    ht5.text_frame.paragraphs[0].font.size = Pt(11)
    ht5.text_frame.paragraphs[0].font.bold = True
    ht5.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    v_tb = s5.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(5.2), Inches(4.5))
    v_tf = v_tb.text_frame
    v_tf.word_wrap = True

    vp1 = v_tf.paragraphs[0]
    vp1.text = "Founder & CTO, SrushtiLabs (srushtilabs.com)"
    vp1.font.size = Pt(12.5)
    vp1.font.bold = True
    vp1.font.color.rgb = COLOR_PRIMARY_DARK
    vp1.space_after = Pt(6)

    vox_body = (
        "• Current Headline: Founder & CTO | Building AI-driven 3D asset platforms and spatial computing solutions @srushtilabs.com\n\n"
        "• Live Application: srushtilabs.com/voxelforge/ — Developing AI-based workflows for generative 3D asset creation and modular bundle assembly.\n\n"
        "• Modular Game-Ready Bundles: Generates optimized, low-poly modular 3D assets ready for immediate assembly in Unreal Engine, Unity, and real-time WebGL engines.\n\n"
        "• Rapid Prototyping Pipeline: Bridges text/image prompts into 3D voxel geometry with automated UV unwrapping and texture generation, cutting 3D asset prototyping time from days to minutes.\n\n"
        "• Lightweight In-Browser Rendering: Specifically engineered to deliver compact asset payloads that stream effortlessly into web browsers without requiring heavy GPU rendering workstations."
    )
    vp2 = v_tf.add_paragraph()
    vp2.text = vox_body
    vp2.font.size = Pt(9.5)
    vp2.font.color.rgb = COLOR_TEXT_MAIN

    # Right: ayam3d & Spatial Engineering Value
    rx5 = Inches(6.68)
    rw5 = Inches(5.85)

    # Card 1: ayam3d
    add_card(s5, rx5, Inches(1.6), rw5, Inches(2.45))
    atb = s5.shapes.add_textbox(rx5 + Inches(0.2), Inches(1.75), rw5 - Inches(0.4), Inches(2.2))
    atf = atb.text_frame
    atf.word_wrap = True

    ap1 = atf.paragraphs[0]
    ap1.text = "ayam3d — Advanced Mesh Synthesis & Retopology"
    ap1.font.size = Pt(12)
    ap1.font.bold = True
    ap1.font.color.rgb = COLOR_PRIMARY_BLUE
    ap1.space_after = Pt(4)

    ap2 = atf.add_paragraph()
    ap2.text = (
        "• Exploratory R&D into automated 3D mesh synthesis, intelligent quad retopology, and PBR texture generation.\n"
        "• Converts complex, unorganized high-poly meshes into clean, lightweight geometry with preserved edge flow.\n"
        "• Solves real-time visualization challenges by creating Level of Detail (LOD) hierarchies for seamless 3D streaming."
    )
    ap2.font.size = Pt(9.5)
    ap2.font.color.rgb = COLOR_TEXT_MAIN

    # Card 2: Industry Relevance
    add_card(s5, rx5, Inches(4.2), rw5, Inches(2.6), bg_color=RGBColor(240, 249, 255), border_color=COLOR_PRIMARY_BLUE)
    itb = s5.shapes.add_textbox(rx5 + Inches(0.2), Inches(4.35), rw5 - Inches(0.4), Inches(2.3))
    itf = itb.text_frame
    itf.word_wrap = True

    ip1 = itf.paragraphs[0]
    ip1.text = "Applied Value to Physical & Digital Infrastructure"
    ip1.font.size = Pt(11.5)
    ip1.font.bold = True
    ip1.font.color.rgb = COLOR_BRAND_DEEP
    ip1.space_after = Pt(4)

    ip2 = itf.add_paragraph()
    ip2.text = (
        "• Bridging Digital Twins with Web Applications: Enables lightweight 3D asset representations of physical spaces, facilities, and structures without multi-gigabyte download bottlenecks.\n"
        "• Practical Spatial Computing: Demonstrates deep technical competence in polygon reduction, coordinate systems, and automated 3D pipelines that directly benefit modern spatial workflows and digital engineering."
    )
    ip2.font.size = Pt(9.5)
    ip2.font.color.rgb = COLOR_TEXT_MAIN

    add_footer(s5, 5, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 6: Visualizer Expert Pedigree (50+ Projects, CAE, Cisco)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, COLOR_BG_LIGHT)
    add_header(s6, "3D Visualization Track Record", "3D Visualizer Expert Pedigree: 50+ Projects, CAE Simulation & Cisco Systems")

    # 3 Column Cards
    c_w = Inches(3.72)
    c_h = Inches(5.2)

    # Col 1: Architectural Visualization
    add_card(s6, Inches(0.8), Inches(1.6), c_w, c_h)
    tb1 = s6.shapes.add_textbox(Inches(0.95), Inches(1.75), c_w - Inches(0.3), c_h - Inches(0.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Architectural 3D & Walkthroughs"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_BLUE
    p.space_after = Pt(2)
    p_sub = tf1.add_paragraph()
    p_sub.text = "35+ Commercial & Residential Projects"
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_BRAND_DEEP
    p_sub.space_after = Pt(8)
    p_desc = tf1.add_paragraph()
    p_desc.text = (
        "• Working with Architect Father: Built foundational expertise interpreting 2D plans, elevations, and structural schematics into accurate 3D models.\n\n"
        "• Photorealistic Visualizations: Produced photorealistic exterior and interior architectural renderings, lighting studies, and texture mappings.\n\n"
        "• 3D Walkthrough Animations: Produced animated camera walkthroughs for real estate developers and commercial builders to visualize spaces prior to construction.\n\n"
        "• CAD & Drafting Precision: Hands-on mastery with CAD drafting tools, layer management, and dimensional accuracy."
    )
    p_desc.font.size = Pt(9)
    p_desc.font.color.rgb = COLOR_TEXT_MAIN

    # Col 2: CAE Simulation Technologies
    add_card(s6, Inches(4.8), Inches(1.6), c_w, c_h)
    tb2 = s6.shapes.add_textbox(Inches(4.95), Inches(1.75), c_w - Inches(0.3), c_h - Inches(0.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "CAE Simulation Technologies"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_BLUE
    p.space_after = Pt(2)
    p_sub = tf2.add_paragraph()
    p_sub.text = "Visual Database Developer (2007 – 2008)"
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_BRAND_DEEP
    p_sub.space_after = Pt(8)
    p_desc = tf2.add_paragraph()
    p_desc.text = (
        "• Flight Simulator Visual Databases: Created high-fidelity 3D model libraries for military and civil aircraft full-flight simulators.\n\n"
        "• Real-Time Terrain Modeling: Built elevation terrain databases, runway markings, airport ground equipment, and navigation landmarks.\n\n"
        "• Strict Polygon & Memory Budgets: Developed real-time 3D models requiring deterministic frame rates (60 FPS locked) with zero visual pop-in or simulation stutter.\n\n"
        "• Sensor & Environmental Simulation: Modeled day/night transitions, weather effects, and infrared/night-vision sensor textures."
    )
    p_desc.font.size = Pt(9)
    p_desc.font.color.rgb = COLOR_TEXT_MAIN

    # Col 3: Cisco Systems Facility Dashboard
    add_card(s6, Inches(8.8), Inches(1.6), c_w, c_h)
    tb3 = s6.shapes.add_textbox(Inches(8.95), Inches(1.75), c_w - Inches(0.3), c_h - Inches(0.3))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Cisco Systems Dashboard"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_BLUE
    p.space_after = Pt(2)
    p_sub = tf3.add_paragraph()
    p_sub.text = "Facility Dashboard & Smart Campus"
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_BRAND_DEEP
    p_sub.space_after = Pt(8)
    p_desc = tf3.add_paragraph()
    p_desc.text = (
        "• Cisco Facility Dashboard (2009 – 2010): Designed and developed an interactive facility dashboard during full-time on-campus consulting at Cisco Bangalore.\n\n"
        "• Physical Infrastructure Telemetry: Visualized data center rack layouts, campus power consumption, cooling metrics, and environmental sensor telemetry.\n\n"
        "• Awarded Cisco Appreciation: Received formal recognition from Cisco leadership in 2010 for delivering intuitive, high-impact data visualization.\n\n"
        "• Continued Cisco Collaboration: Managed multiple Angular dashboards and UI applications for Cisco Systems via UST Global (2015 – 2019)."
    )
    p_desc.font.size = Pt(9)
    p_desc.font.color.rgb = COLOR_TEXT_MAIN

    add_footer(s6, 6, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 7: Role & Core Responsibilities Across Career
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, COLOR_BG_LIGHT)
    add_header(s7, "Leadership & Execution", "Core Responsibilities & Technical Leadership Across 16+ Years")

    grid_w = Inches(5.72)
    grid_h = Inches(2.45)
    cols = [Inches(0.8), Inches(6.8)]
    rows = [Inches(1.6), Inches(4.3)]

    pillars_data = [
        ("1. Scalable Architecture & System Design",
         COLOR_PRIMARY_BLUE,
         [
             "Architected enterprise frontend applications using React, Angular, and LitElement web components.",
             "Designed cross-platform UI Design Systems adopted across diverse development teams (Moonraft / TLC Hotels).",
             "Established robust contract boundaries using REST, WebSockets, and event-driven data streaming.",
             "Evaluated and selected modern cloud and web tech stacks for rapid, scalable client product delivery."
         ]),
        ("2. Operational Resilience & High Reliability",
         COLOR_BRAND_DEEP,
         [
             "Architected zero-loss offline-first caching mechanisms (IndexedDB/Service Workers) for field use (Twitan).",
             "Engineered life-critical clinical UI rendering with zero memory leaks for 24/7 hospital ICU units (Philips ICCA).",
             "Supervised cloud deployments on GCP, Firebase, and AWS with automated scaling and log monitoring.",
             "Designed real-time fallback pipelines ensuring continuous operational availability during network drops."
         ]),
        ("3. Spatial Computing & 3D Asset Pipelines",
         COLOR_ACCENT_AMBER,
         [
             "Developed AI-assisted workflows (Voxelforge AI) for modular low-poly 3D asset generation (Unreal / Unity / WebGL).",
             "Directed R&D in automated 3D mesh retopology and PBR texture generation for lightweight web streaming (ayam3d).",
             "Created real-time 3D flight simulator databases with strict polygon and frame rate budgets (CAE).",
             "Built interactive 3D facility and infrastructure monitoring dashboards (Cisco Systems)."
         ]),
        ("4. Agile Leadership, Mentorship & Delivery",
         COLOR_SUCCESS,
         [
             "Led, mentored, and inspired cross-functional engineering squads of developers, UX designers, and QA analysts.",
             "Instituted design-to-code pipelines involving users early through rapid prototypes and usability feedback.",
             "Managed client deliverables, stakeholder communication, sprint planning, and engineering velocity.",
             "Recipient of multiple corporate excellence awards at UST Global and Cisco for leadership and client commitment."
         ])
    ]

    for idx, (p_title, p_color, p_bullets) in enumerate(pillars_data):
        c_x = cols[idx % 2]
        c_y = rows[idx // 2]
        add_card(s7, c_x, c_y, grid_w, grid_h)

        hbar = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x + Inches(0.01), c_y + Inches(0.01), grid_w - Inches(0.02), Inches(0.42))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = p_color
        hbar.line.fill.background()

        htext = s7.shapes.add_textbox(c_x + Inches(0.2), c_y + Inches(0.06), grid_w - Inches(0.4), Inches(0.35))
        htf = htext.text_frame
        hp = htf.paragraphs[0]
        hp.text = p_title
        hp.font.size = Pt(11)
        hp.font.bold = True
        hp.font.color.rgb = RGBColor(255, 255, 255)

        btb = s7.shapes.add_textbox(c_x + Inches(0.2), c_y + Inches(0.48), grid_w - Inches(0.4), grid_h - Inches(0.55))
        btf = btb.text_frame
        btf.word_wrap = True
        for b_i, bullet in enumerate(p_bullets):
            bp = btf.paragraphs[0] if b_i == 0 else btf.add_paragraph()
            bp.text = f"• {bullet}"
            bp.font.size = Pt(9.5)
            bp.font.color.rgb = COLOR_TEXT_MAIN
            bp.space_after = Pt(3)

    add_footer(s7, 7, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 8: Major Achievements, Milestones & Awards
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, COLOR_BG_LIGHT)
    add_header(s8, "Milestones & Recognition", "Major Achievements, Proven Track Record & Formal Acknowledgements")

    # 4 Authentic Metric Cards
    metrics = [
        ("16+ Years", "Professional Track Record", "Extensive leadership across enterprise software, healthcare, startups & 3D tech"),
        ("50+ Projects", "3D & Visualization Portfolio", "Delivered across architectural visualization, CAD, flight simulation & spatial AI"),
        ("4 Formal Awards", "Industry Acknowledgements", "Cisco Facility Appreciation + 3 UST Global Excellence & Client-First awards"),
        ("100% Offline", "Data Integrity at Edge", "Zero data loss during field network drops with transactional replay queues")
    ]
    m_w = Inches(2.78)
    m_h = Inches(1.5)
    for i, (stat, label, detail) in enumerate(metrics):
        mx = Inches(0.8) + i * (m_w + Inches(0.2))
        add_card(s8, mx, Inches(1.6), m_w, m_h, bg_color=RGBColor(255, 255, 255), border_color=COLOR_PRIMARY_BLUE)

        tb = s8.shapes.add_textbox(mx + Inches(0.12), Inches(1.72), m_w - Inches(0.24), m_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = stat
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_BLUE
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_PRIMARY_DARK
        p2.space_after = Pt(2)

        p3 = tf.add_paragraph()
        p3.text = detail
        p3.font.size = Pt(8.5)
        p3.font.color.rgb = COLOR_TEXT_MUTED

    # Detailed Career Achievements List
    achievements = [
        ("Philips Intellispace Critical Care & Anesthesia (Cyient)",
         "Led the frontend technical development of clinical modules for Philips Healthcare, delivering life-critical patient monitoring with rock-solid stability, zero memory leaks, and full healthcare regulatory compliance."),
        ("Startup Co-Founder & Director (Technoyana & Twitan)",
         "Successfully co-founded and steered Technoyana Digital Transformation Services and Twitan.com; built and shipped commercial fintech apps, product search engines, and the Shutlify badminton tournament operating system."),
        ("Cross-Platform UI Design System (Moonraft Innovation Labs)",
         "Architected a cross-platform UI design system using the LitElement web components library, creating a reusable component standard leveraged seamlessly across React, Angular, and Polymer teams for TLC Hotels."),
        ("Formal Corporate Awards & Cisco Recognition",
         "Received Cisco Systems Appreciation for developing the campus Facility Dashboard (2010); recognized with 3 Certificates of Appreciation at UST Global ('Inspiring People' 2018, 'Putting Client First' 2016, 'Living the Values' 2015).")
    ]

    for i, (title, body) in enumerate(achievements):
        ay = Inches(3.25) + i * Inches(0.92)
        add_card(s8, Inches(0.8), ay, Inches(11.733), Inches(0.84))
        atb = s8.shapes.add_textbox(Inches(1.0), ay + Inches(0.08), Inches(11.3), Inches(0.68))
        atf = atb.text_frame
        atf.word_wrap = True

        ap1 = atf.paragraphs[0]
        ap1.text = f"★  {title}"
        ap1.font.size = Pt(11)
        ap1.font.bold = True
        ap1.font.color.rgb = COLOR_PRIMARY_DARK
        ap1.space_after = Pt(2)

        ap2 = atf.add_paragraph()
        ap2.text = body
        ap2.font.size = Pt(9)
        ap2.font.color.rgb = COLOR_TEXT_MAIN

    add_footer(s8, 8, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 9: Challenges Faced & Pragmatic Resolutions
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, COLOR_BG_LIGHT)
    add_header(s9, "Problem Solving & Engineering Rigor", "Real-World Technical Challenges Faced & Pragmatic Resolutions")

    c_cards = [
        ("Challenge 1: Unstable Network Connectivity in Physical Venues (Twitan / Technoyana)",
         "Operating live tournament arbitration in venues with crowded RF environments, intermittent Wi-Fi drops, and mobile dead-zones caused transaction failures and loss of live scores.",
         "Resolution Implemented:",
         "Architected an 'Offline-First' state engine using IndexedDB client persistence and service workers. Implemented optimistic UI updates and a transactional event replay queue that automatically resynchronizes once the network reconnects.",
         "Impact: 100% data preservation, zero operator workflow interruptions, seamless user experience."),

        ("Challenge 2: Clinical UI Performance & Memory Leaks in 24/7 ICU Monitoring (Philips ICCA)",
         "Bedside patient monitors run uninterrupted for weeks. Complex real-time streaming vitals and continuous DOM updates risked browser memory accumulation and UI lag during critical clinical events.",
         "Resolution Implemented:",
         "Enforced strict immutable state updates, decoupled rendering loops from data ingestion, implemented virtualized list rendering, and utilized Chrome DevTools heap profiling to systematically eliminate closure leaks.",
         "Impact: Guaranteed stable 60 FPS UI rendering and flat memory profiles across multi-day continuous ICU sessions."),

        ("Challenge 3: High 3D Asset Payload & Browser Lag in Web-Based 3D (SrushtiLabs / Voxelforge AI)",
         "Delivering detailed 3D models over the web created large download sizes, long initialization times, and frame drops on lower-spec client laptops and mobile devices.",
         "Resolution Implemented:",
         "Developed automated retopology workflows (ayam3d) to decimate unnecessary polygons while preserving silhouette geometry; baked normal maps and bundled assets into lightweight modular formats for WebGL.",
         "Impact: Reduced 3D model payload sizes by up to 75%, enabling instant in-browser loading and smooth real-time manipulation.")
    ]

    for i, (ch_title, ch_desc, res_lbl, res_desc, impact) in enumerate(c_cards):
        cy = Inches(1.6) + i * Inches(1.76)
        add_card(s9, Inches(0.8), cy, Inches(11.733), Inches(1.62))

        ltb = s9.shapes.add_textbox(Inches(1.0), cy + Inches(0.12), Inches(4.8), Inches(1.4))
        ltf = ltb.text_frame
        ltf.word_wrap = True
        lp1 = ltf.paragraphs[0]
        lp1.text = ch_title
        lp1.font.size = Pt(11)
        lp1.font.bold = True
        lp1.font.color.rgb = COLOR_BRAND_DEEP
        lp1.space_after = Pt(3)
        lp2 = ltf.add_paragraph()
        lp2.text = ch_desc
        lp2.font.size = Pt(9)
        lp2.font.color.rgb = COLOR_TEXT_MAIN

        vdiv = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.0), cy + Inches(0.15), Inches(0.02), Inches(1.3))
        vdiv.fill.solid()
        vdiv.fill.fore_color.rgb = COLOR_CARD_BORDER
        vdiv.line.fill.background()

        rtb = s9.shapes.add_textbox(Inches(6.2), cy + Inches(0.12), Inches(6.1), Inches(1.4))
        rtf = rtb.text_frame
        rtf.word_wrap = True
        rp1 = rtf.paragraphs[0]
        rp1.text = res_lbl
        rp1.font.size = Pt(11)
        rp1.font.bold = True
        rp1.font.color.rgb = COLOR_PRIMARY_BLUE
        rp1.space_after = Pt(2)
        rp2 = rtf.add_paragraph()
        rp2.text = res_desc
        rp2.font.size = Pt(9)
        rp2.font.color.rgb = COLOR_TEXT_MAIN
        rp2.space_after = Pt(3)
        rp3 = rtf.add_paragraph()
        rp3.text = f"Result: {impact}"
        rp3.font.size = Pt(9)
        rp3.font.bold = True
        rp3.font.color.rgb = COLOR_SUCCESS

    add_footer(s9, 9, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 10: Lessons Learned & Engineering Best Practices
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, COLOR_BG_LIGHT)
    add_header(s10, "Engineering Philosophy", "Lessons Learned & Core Engineering Best Practices")

    best_practices = [
        ("1. Involve Users Early Through Rapid Prototypes & Iterative Feedback",
         "From my days establishing design-to-code pipelines at ThoughtFocus to architecting clinical modules at Cyient, building early interactive prototypes with real users uncovers fundamental workflow flaws before expensive engineering code is written."),
        ("2. Design for Real-World Edge Conditions, Not Ideal Lab Environments",
         "Software rarely operates in perfect network and hardware conditions. Building applications with offline caching, graceful degradation, and asynchronous retry logic ensures systems remain reliable whether on a hospital floor, sports arena, or remote field office."),
        ("3. Standardize Design Systems & Component Libraries to Prevent Technical Debt",
         "Re-inventing UI components across teams leads to inconsistent user experiences and bloated maintenance. Investing upfront in unified, framework-agnostic design systems (as demonstrated at Moonraft with LitElement) dramatically accelerates organizational velocity."),
        ("4. Ground Technology Decisions in Concrete Business & Human Impact",
         "Technology is a vehicle for solving human problems. Whether building spatial AI tools (Voxelforge AI), sports arbitration engines (Twitan), or ICU monitoring (Philips), the ultimate metric of success is operational simplicity, reliability, and tangible value to the user.")
    ]

    for i, (title, desc) in enumerate(best_practices):
        by = Inches(1.6) + i * Inches(1.3)
        add_card(s10, Inches(0.8), by, Inches(11.733), Inches(1.15))

        tag = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), by, Inches(0.12), Inches(1.15))
        tag.fill.solid()
        tag.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
        tag.line.fill.background()

        btb = s10.shapes.add_textbox(Inches(1.15), by + Inches(0.12), Inches(11.1), Inches(0.95))
        btf = btb.text_frame
        btf.word_wrap = True

        bp1 = btf.paragraphs[0]
        bp1.text = title
        bp1.font.size = Pt(12)
        bp1.font.bold = True
        bp1.font.color.rgb = COLOR_PRIMARY_DARK
        bp1.space_after = Pt(4)

        bp2 = btf.add_paragraph()
        bp2.text = desc
        bp2.font.size = Pt(9.5)
        bp2.font.color.rgb = COLOR_TEXT_MAIN

    add_footer(s10, 10, TOTAL_SLIDES)

    # ==========================================
    # SLIDE 11: Suggestions for Future Process Improvements & Innovation
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, COLOR_PRIMARY_DARK)

    # Center Card
    c_w = Inches(11.733)
    c_h = Inches(5.2)
    cx = Inches(0.8)
    cy = Inches(1.5)
    add_card(s11, cx, cy, c_w, c_h, bg_color=RGBColor(24, 34, 53), border_color=RGBColor(51, 65, 85))

    q_tb = s11.shapes.add_textbox(cx + Inches(0.5), cy + Inches(0.35), c_w - Inches(1.0), c_h - Inches(0.7))
    q_tf = q_tb.text_frame
    q_tf.word_wrap = True

    qp1 = q_tf.paragraphs[0]
    qp1.text = "VISION FOR FUTURE PROCESS IMPROVEMENTS & SCALE"
    qp1.font.size = Pt(11.5)
    qp1.font.bold = True
    qp1.font.color.rgb = COLOR_CYAN_ACCENT
    qp1.space_after = Pt(4)

    qp2 = q_tf.add_paragraph()
    qp2.text = "Leveraging Proven Expertise in Scalable Systems, UX & Spatial Tech"
    qp2.font.size = Pt(21)
    qp2.font.bold = True
    qp2.font.color.rgb = RGBColor(255, 255, 255)
    qp2.space_after = Pt(12)

    proposals = [
        ("Unified Spatial & Web Integration", "Bridge complex 3D CAD, architectural models, and spatial data into lightweight, interactive web formats using progressive LOD and automated retopology (leveraging Voxelforge AI / ayam3d principles)."),
        ("Resilient Edge-First Mobile Workflows", "Implement offline-first client persistence and background transactional queuing across field-facing mobile applications, ensuring zero work loss regardless of remote connectivity challenges."),
        ("Standardized Enterprise Design Systems", "Establish centralized, cross-platform component libraries and design tokens to streamline multi-team digital product delivery, enforce visual consistency, and eliminate duplicated engineering effort."),
        ("AI-Assisted Prototyping & Development", "Incorporate generative AI and automation tools into early product discovery and asset creation pipelines, cutting design-to-code iteration cycles by 40%.")
    ]

    for p_title, p_desc in proposals:
        pp1 = q_tf.add_paragraph()
        pp1.text = f"• {p_title}: "
        pp1.font.bold = True
        pp1.font.size = Pt(10.5)
        pp1.font.color.rgb = COLOR_CYAN_ACCENT
        pp2 = q_tf.add_paragraph()
        pp2.text = f"  {p_desc}"
        pp2.font.size = Pt(9.5)
        pp2.font.color.rgb = RGBColor(203, 213, 225)
        pp2.space_after = Pt(6)

    qp_contact = q_tf.add_paragraph()
    qp_contact.text = "\nMaheshchandra Hegde  |  hid.mahesh@gmail.com  |  +91 9535253329 / 7022407280  |  hegdemahesh.in"
    qp_contact.font.size = Pt(11)
    qp_contact.font.bold = True
    qp_contact.font.color.rgb = RGBColor(255, 255, 255)

    add_footer(s11, 11, TOTAL_SLIDES, dark=True)

    prs.save(output_pptx_path)
    print(f"Successfully generated Authentic 11-Slide PowerPoint presentation at: {output_pptx_path}")

if __name__ == "__main__":
    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_file = os.path.join(out_dir, "Obayashi_Senior_Manager_IT_Presentation_Maheshchandra_Hegde.pptx")
    create_deck(out_file)
